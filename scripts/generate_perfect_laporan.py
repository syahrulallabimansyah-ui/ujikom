# -*- coding: utf-8 -*-
"""
Script to generate LAPORAN_UKK_AKSA_NOVA.md and LAPORAN_UKK_AKSA_NOVA.docx
Incorporating all corrections:
1. Cover formal, Lembar Pengesahan, Kata Pengantar, Daftar Isi, Daftar Gambar, Daftar Tabel
2. Sub-bab 2.3 Jadwal & Alur Pengerjaan Waterfall (Tabel)
3. Sinkronisasi Daftar Isi & Judul (3.2 User Persona & Aktor, 3.3 24 Spesifikasi FR-01 s/d FR-24)
4. Heading 1/2/3/4 Word resmi sehingga Navigation Pane Word aktif
5. Page Number di footer Word
6. Bersihkan semua simbol LaTeX (ganti dengan unicode ->, x, <=, >=, etc.)
7. Hasil implementasi UI (Screenshots 5.2) & potongan kode penting di BAB V
8. Wireframe / Mockup di 4.6
9. Koreksi detail: UC-06 Tamu, UC-01 include UC-02, 1 Class Diagram 14 kelas, 18 Black Box test cases
10. Nama jurusan resmi: PENGEMBANGAN PERANGKAT LUNAK DAN GIM (PPLG)
11. Daftar Pustaka formal
12. Lampiran katalog diagram pelengkap
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from PIL import Image

def clean_latex_symbols(text):
    text = text.replace('$\\rightarrow$', '→')
    text = text.replace('$\\dashrightarrow$', '⇢')
    text = text.replace('$\\times$', '×')
    text = text.replace('$\\le 768px$', '≤ 768px')
    text = text.replace('$\\le 5$', '≤ 5')
    text = text.replace('$\\le 1$', '≤ 1')
    text = text.replace('$\\ge 360 \\times 640$', '≥ 360 × 640')
    text = text.replace('$\\blacklozenge$', '◆')
    text = text.replace('\\ll\\text{create}\\gg', '«create»')
    text = text.replace('\\ll\\text{use}\\gg', '«use»')
    text = text.replace('$\\le$', '≤')
    text = text.replace('$\\ge$', '≥')
    return text

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_page_number_to_section(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Halaman ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(120, 120, 120)
    
    # Add page field
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p._element.append(fldSimple)
    
    run2 = p.add_run(" | Laporan UKK - AKSA NOVA")
    run2.font.name = "Calibri"
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(140, 140, 140)

print("Setup helper functions completed.")
